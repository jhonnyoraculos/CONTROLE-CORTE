"""Inspection and parsing of the historical cut-control spreadsheet.

This module deliberately does not write to the database or reconcile rows with
the current service export. The historical sheet has inconsistent identifiers
and a numeric header whose meaning requires an explicit user decision.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from importers.normalizers import normalize_identifier, parse_date, parse_decimal


@dataclass
class LegacyRow:
    order_code: str
    service_code: str | None
    fields: dict[str, object]
    raw_data: dict[str, object]
    row_number: int


@dataclass
class LegacyResult:
    rows: list[LegacyRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def preview(self) -> list[dict[str, object]]:
        """First 20 parsed rows in a JSON-safe shape for the review screen."""
        return [
            {
                "row_number": row.row_number,
                "order_code": row.order_code,
                "service_code": row.service_code,
                "fields": {key: _plain(value) for key, value in row.fields.items()},
                "raw_data": row.raw_data,
            }
            for row in self.rows[:20]
        ]


_HEADER_TO_FIELD = {
    "data pedido": "data_pedido",
    "vendedor": "vendedor_legacy",
    "loja venda": "loja_venda",
    "no pedido": "numero_pedido",
    "n pedido": "numero_pedido",
    "numero pedido": "numero_pedido",
    "id cliente": "id_cliente_legacy",
    "data carregamento": "data_carregamento",
    "modalidade entrega retira": "modalidade",
    "remessa": "remessa",
    "cidade": "cidade",
    "endereco de entrega": "endereco",
    "bairro": "bairro",
    "peso": "peso",
    "total": "valor_total",
    "tipo ordem de venda": "tipo_ordem_venda",
    "central producao": "central_producao",
    "cod plano projeto": "codigo_plano_projeto",
    "tipo de plano": "tipo_plano",
    "qtd chapas": "chapas",
    "qtd cortes": "cortes",
    "qtd mt fita": "metros_fita_legacy",
    "qtd usinagem furacao": "usinagens",
    "data producao": "data_producao",
    "turno": "turno",
    "capacidade producao": "capacidade_producao_legacy",
    "status": "status_legacy",
    "seccionadora": "seccionadora",
    "data finalizacao servico": "data_finalizacao_servico",
    "observacoes": "observacoes_legacy",
    "tempo de producao": "tempo_producao_legacy",
}
_REQUIRED = {"numero_pedido", "data_pedido", "data_carregamento", "status_legacy"}
_DECIMAL_FIELDS = {
    "peso", "valor_total", "chapas", "cortes", "metros_fita_legacy",
    "usinagens", "tempo_producao_legacy",
}
_DATE_FIELDS = {
    "data_pedido", "data_carregamento", "data_producao", "data_finalizacao_servico",
}
_IDENTIFIER_FIELDS = {"id_cliente_legacy", "remessa", "codigo_plano_projeto"}
_MAPPING_VALUES = {"cliente_nome", "ignorar"}


def _key(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text.casefold())).strip()


def _plain(value: Any) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _find_sheet(workbook: Any) -> tuple[Any, int, dict[str, int], dict[int, Any]] | None:
    options = []
    for sheet in workbook.worksheets:
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25)), 1
        ):
            mapping: dict[str, int] = {}
            headers: dict[int, Any] = {}
            for column, cell in enumerate(cells, 1):
                if cell.value is not None:
                    headers[column] = cell.value
                field_name = _HEADER_TO_FIELD.get(_key(cell.value))
                if field_name and field_name not in mapping:
                    mapping[field_name] = column
            if _REQUIRED.issubset(mapping) and len(mapping) >= 12:
                options.append((len(mapping), sheet, row_number, mapping, headers))
    if not options:
        return None
    _, sheet, header_row, mapping, headers = max(options, key=lambda option: option[0])
    return sheet, header_row, mapping, headers


def _load(data: bytes) -> tuple[Any, tuple[Any, int, dict[str, int], dict[int, Any]]]:
    if not data:
        raise ValueError("Arquivo Excel vazio.")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError(f"Arquivo Excel inválido: {exc}") from exc
    found = _find_sheet(workbook)
    if found is None:
        workbook.close()
        raise ValueError("Cabeçalhos da planilha histórica não encontrados.")
    return workbook, found


def inspect_legacy(data: bytes) -> list[dict[str, object]]:
    """Describe source columns and show samples before the user maps header 27.

    The numeric header is deliberately returned without a suggested field.
    """
    workbook, (sheet, header_row, mapping, headers) = _load(data)
    try:
        field_by_column = {column: name for name, column in mapping.items()}
        inspected = []
        for column in range(1, sheet.max_column + 1):
            original_header = headers.get(column)
            values = []
            for cells in sheet.iter_rows(
                min_row=header_row + 1,
                max_row=min(sheet.max_row, header_row + 100),
                min_col=column,
                max_col=column,
            ):
                value = cells[0].value
                if value is not None and str(value).strip() and _plain(value) not in values:
                    values.append(_plain(value))
                    if len(values) == 3:
                        break
            ambiguous = _key(original_header) == "27"
            inspected.append({
                "column": get_column_letter(column),
                "header": _plain(original_header),
                "suggested_field": None if ambiguous else field_by_column.get(column),
                "requires_mapping": ambiguous,
                "sample_values": values,
            })
        return inspected
    finally:
        workbook.close()


def parse_legacy(data: bytes, mapping: dict[str, str]) -> LegacyResult:
    """Parse the historical sheet after an explicit mapping decision for 27.

    Mapping keys are Excel letters (``F``) or the literal header ``27``.
    Accepted decisions for that column are ``cliente_nome`` and ``ignorar``.
    This function does not merge or persist legacy records.
    """
    result = LegacyResult()
    try:
        workbook, (sheet, header_row, columns, headers) = _load(data)
    except ValueError as exc:
        result.errors.append({"row": None, "field": None, "message": str(exc)})
        return result
    try:
        if not isinstance(mapping, dict):
            result.errors.append({
                "row": header_row, "field": "mapping", "message": "Mapeamento manual obrigatório.",
            })
            return result
        ambiguous_columns = [column for column, header in headers.items() if _key(header) == "27"]
        if len(ambiguous_columns) > 1:
            result.errors.append({
                "row": header_row, "field": "mapping", "message": "Há mais de uma coluna com cabeçalho 27; revise o arquivo.",
            })
            return result
        for column in ambiguous_columns:
            letter = get_column_letter(column)
            choices = [mapping[key] for key in (letter, letter.lower(), "27") if key in mapping]
            if not choices or len(set(choices)) != 1 or choices[0] not in _MAPPING_VALUES:
                result.errors.append({
                    "row": header_row,
                    "field": letter,
                    "message": "Confirme manualmente a coluna com cabeçalho 27: mapping={'" + letter + "': 'cliente_nome'} ou 'ignorar'.",
                    "value": _plain(headers[column]),
                })
                return result
            if choices[0] == "cliente_nome":
                columns["cliente_nome"] = column
            else:
                result.warnings.append(f"Coluna {letter} (cabeçalho 27) ignorada por escolha explícita.")

        if not ambiguous_columns:
            result.warnings.append("Nenhuma coluna com cabeçalho 27 encontrada neste arquivo histórico.")
        header_labels = {
            col: str(_plain(headers.get(col))) if headers.get(col) is not None else get_column_letter(col)
            for col in range(1, sheet.max_column + 1)
        }
        skipped_no_order = 0
        text_plan_codes = 0
        for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if all(cell.value is None or cell.value == "" for cell in cells):
                continue
            raw_data = {
                header_labels[col]: _plain(cell.value)
                for col, cell in enumerate(cells, 1)
                if col in header_labels
            }
            values = {
                name: cells[column - 1].value if column <= len(cells) else None
                for name, column in columns.items()
            }
            try:
                order_code = normalize_identifier(values.get("numero_pedido"))
            except ValueError as exc:
                result.errors.append({
                    "row": row_number, "field": "numero_pedido", "message": str(exc),
                    "value": _plain(values.get("numero_pedido")),
                })
                continue
            if not order_code:
                skipped_no_order += 1
                continue
            if len(order_code) > 80:
                result.errors.append({
                    "row": row_number, "field": "numero_pedido",
                    "message": "Número do pedido excede 80 caracteres.",
                })
                continue
            if not re.fullmatch(r"\d+", order_code):
                raw_data["_needs_order_review"] = True
                result.warnings.append(
                    f"Linha {row_number}: número do pedido {order_code!r} não é simples; revisar vínculo."
                )

            fields: dict[str, object] = {}
            for name, value in values.items():
                if name == "numero_pedido":
                    continue
                if name in _DECIMAL_FIELDS:
                    try:
                        fields[name] = parse_decimal(value)
                    except ValueError as exc:
                        fields[name] = None
                        result.errors.append({
                            "row": row_number, "field": name, "message": str(exc), "value": _plain(value),
                        })
                elif name in _DATE_FIELDS:
                    date_warnings: list[str] = []
                    try:
                        parsed = parse_date(value, warnings=date_warnings)
                        fields[name] = parsed.date() if isinstance(parsed, datetime) else parsed
                    except ValueError as exc:
                        fields[name] = None
                        result.errors.append({
                            "row": row_number, "field": name, "message": str(exc), "value": _plain(value),
                        })
                    result.warnings.extend(
                        f"Linha {row_number}, {name}: {message}" for message in date_warnings
                    )
                elif name in _IDENTIFIER_FIELDS:
                    try:
                        fields[name] = normalize_identifier(value) or None
                    except ValueError as exc:
                        fields[name] = None
                        result.errors.append({
                            "row": row_number, "field": name, "message": str(exc), "value": _plain(value),
                        })
                else:
                    fields[name] = _text(value)
            plan_code = fields.get("codigo_plano_projeto")
            service_code = plan_code if isinstance(plan_code, str) and re.fullmatch(r"\d+", plan_code) else None
            if plan_code and service_code is None:
                text_plan_codes += 1
            result.rows.append(LegacyRow(order_code, service_code, fields, raw_data, row_number))
        if skipped_no_order:
            result.warnings.append(f"{skipped_no_order} linha(s) sem número de pedido foram ignoradas.")
        if text_plan_codes:
            result.warnings.append(
                f"{text_plan_codes} código(s) de plano/projeto não numéricos não foram usados como Serviço."
            )
        return result
    finally:
        workbook.close()
