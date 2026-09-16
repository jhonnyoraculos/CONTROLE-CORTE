"""Read the current CorteCloud service export without changing stored data."""

from __future__ import annotations

import re
import os
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from importers.normalizers import normalize_identifier, parse_date, parse_decimal


@dataclass
class ExcelRow:
    order_code: str
    service_code: str
    fields: dict[str, object]
    raw_data: dict[str, object]
    row_number: int


@dataclass
class ExcelResult:
    rows: list[ExcelRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


_HEADERS = {
    "servico": "servico",
    "codigo interno": "codigo_interno",
    "cliente": "cliente_origem",
    "cliente final": "cliente_final_origem",
    "vendedor": "vendedor_origem",
    "enviado producao": "enviado_producao",
    "central": "central",
    "linha producao": "linha_producao",
    "chapas": "chapas",
    "cortes": "cortes",
    "metros lineares de corte": "metros_lineares_corte",
    "pecas": "pecas",
    "fita aplicada": "fita_aplicada",
    "usinagens": "usinagens",
    "corte realizado em": "corte_realizado_em",
    "fitagem realizada em": "fitagem_realizada_em",
    "usinagem realizada em": "usinagem_realizada_em",
    "previsao entrega": "previsao_entrega",
    "observacao": "observacao_origem",
}
_REQUIRED = {"servico", "codigo_interno"}
_TEXT_FIELDS = (
    "cliente_origem", "cliente_final_origem", "vendedor_origem", "central",
    "linha_producao", "observacao_origem",
)
_DECIMAL_FIELDS = (
    "chapas", "cortes", "metros_lineares_corte", "pecas", "fita_aplicada", "usinagens",
)
_DATE_FIELDS = (
    "enviado_producao", "corte_realizado_em", "fitagem_realizada_em",
    "usinagem_realizada_em", "previsao_entrega",
)


def _header_key(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).strip())
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text.casefold())).strip()


def _plain(value: Any) -> object:
    """Make a workbook value safe for JSON storage while retaining its content."""
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    return result or None


def _find_sheet(workbook: Any) -> tuple[Any, int, dict[str, int]] | None:
    """Identify the export by its header structure, never by sheet name."""
    candidates = []
    for sheet in workbook.worksheets:
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25)), 1
        ):
            mapping: dict[str, int] = {}
            for cell in cells:
                field_name = _HEADERS.get(_header_key(cell.value))
                if field_name and field_name not in mapping:
                    mapping[field_name] = cell.column
            if _REQUIRED.issubset(mapping) and len(mapping) >= 4:
                candidates.append((len(mapping), sheet, row_number, mapping))
    if not candidates:
        return None
    _, sheet, header_row, mapping = max(candidates, key=lambda candidate: candidate[0])
    return sheet, header_row, mapping


def parse_excel(data: bytes) -> ExcelResult:
    """Parse one services workbook for preview and later transactional import.

    Blank cells are explicit values for columns present in the workbook.
    Missing columns and unreadable optional values are omitted so they cannot
    erase an existing value during a partial or faulty import.
    """
    result = ExcelResult()
    if not data:
        result.errors.append({"row": None, "field": None, "message": "Arquivo Excel vazio."})
        return result
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        result.errors.append({"row": None, "field": None, "message": f"Arquivo Excel inválido: {exc}"})
        return result
    try:
        found = _find_sheet(workbook)
        if found is None:
            result.errors.append({
                "row": None,
                "field": None,
                "message": "Cabeçalhos da planilha de serviços não encontrados. Esperados Serviço, Código Interno e ao menos duas colunas de serviço reconhecidas.",
            })
            return result
        sheet, header_row, columns = found
        headers = {
            column: _text(sheet.cell(header_row, column).value) or get_column_letter(column)
            for column in range(1, sheet.max_column + 1)
        }
        missing_optional = set(_HEADERS.values()) - set(columns)
        if missing_optional:
            result.warnings.append(
                "Colunas opcionais ausentes: " + ", ".join(sorted(missing_optional)) + "."
            )
        seen_services: set[tuple[str, str]] = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if all(cell.value is None or cell.value == "" for cell in cells):
                continue
            raw_data = {
                headers[column]: _plain(cell.value)
                for column, cell in enumerate(cells, 1)
                if column in headers
            }
            row_values = {
                name: cells[column - 1].value if column <= len(cells) else None
                for name, column in columns.items()
            }
            formula_fields = {
                name for name, column in columns.items()
                if column <= len(cells) and cells[column - 1].data_type == "f"
            }
            for name in sorted(formula_fields):
                result.errors.append({
                    "row": row_number, "field": name,
                    "message": "Fórmula em campo de importação; informe um valor fixo.",
                    "value": _plain(row_values[name]),
                })
                row_values[name] = None

            try:
                order_code = normalize_identifier(row_values.get("codigo_interno"))
                service_code = normalize_identifier(row_values.get("servico"))
            except ValueError as exc:
                result.errors.append({"row": row_number, "field": "identificador", "message": str(exc)})
                continue
            if not order_code or not service_code:
                result.errors.append({
                    "row": row_number, "field": "identificador",
                    "message": "Serviço e Código Interno são obrigatórios.",
                    "value": {"Serviço": _plain(row_values.get("servico")), "Código Interno": _plain(row_values.get("codigo_interno"))},
                })
                continue
            if len(order_code) > 80 or len(service_code) > 80:
                result.errors.append({
                    "row": row_number, "field": "identificador",
                    "message": "Identificador excede 80 caracteres.",
                })
                continue
            if re.fullmatch(r"\d+(?:\s+\d+)+", order_code):
                raw_data["_needs_order_review"] = True
                result.warnings.append(
                    f"Linha {row_number}: Código Interno contém múltiplos números ({order_code}); revisar vínculo antes de importar."
                )
            pair = order_code, service_code
            if pair in seen_services:
                result.warnings.append(f"Linha {row_number}: serviço {service_code} repetido para pedido {order_code} no arquivo.")
            seen_services.add(pair)

            fields: dict[str, object] = {
                "codigo_servico": service_code,
                "codigo_servico_normalizado": service_code,
            }
            for name in _TEXT_FIELDS:
                if name in columns and name not in formula_fields:
                    fields[name] = _text(row_values.get(name))
            for name in _DECIMAL_FIELDS:
                if name not in columns or name in formula_fields:
                    continue
                try:
                    fields[name] = parse_decimal(row_values.get(name))
                except ValueError as exc:
                    result.errors.append({
                        "row": row_number, "field": name, "message": str(exc),
                        "value": _plain(row_values.get(name)),
                    })
            for name in _DATE_FIELDS:
                if name not in columns or name in formula_fields:
                    continue
                date_warnings: list[str] = []
                try:
                    parsed = parse_date(row_values.get(name), warnings=date_warnings)
                    if name == "previsao_entrega":
                        fields[name] = parsed.date() if isinstance(parsed, datetime) else parsed
                    elif isinstance(parsed, date) and not isinstance(parsed, datetime):
                        fields[name] = datetime.combine(parsed, time.min).replace(
                            tzinfo=ZoneInfo(os.getenv("APP_TIMEZONE", "America/Sao_Paulo")))
                    else:
                        fields[name] = (parsed.replace(
                            tzinfo=ZoneInfo(os.getenv("APP_TIMEZONE", "America/Sao_Paulo")))
                            if isinstance(parsed, datetime) and parsed.tzinfo is None else parsed)
                except ValueError as exc:
                    result.errors.append({
                        "row": row_number, "field": name, "message": str(exc),
                        "value": _plain(row_values.get(name)),
                    })
                result.warnings.extend(f"Linha {row_number}, {name}: {message}" for message in date_warnings)
            result.rows.append(ExcelRow(order_code, service_code, fields, raw_data, row_number))
        if not result.rows and not result.errors:
            result.warnings.append("Planilha de serviços sem linhas de dados.")
        return result
    finally:
        workbook.close()
