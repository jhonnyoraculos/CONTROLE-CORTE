"""End-to-end import checks using the documents supplied for this project."""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
import pymupdf
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from db.models import Audit, Base, Document, ImportError, Order, OrderItem, PendingLink, Service
from services.importacao_service import ImportService


EXCEL_PATH = Path(r"C:\Users\JRFerragens\Downloads\PLANILHA (1).xlsx")
PDF_PATH = Path(
    r"C:\Users\JRFerragens\OneDrive - JR Ferragens & Madeiras"
    r"\Área de Trabalho\CARRINHOFIN.pdf"
)


@pytest.fixture(scope="module")
def real_files() -> tuple[bytes, bytes]:
    if not EXCEL_PATH.is_file() or not PDF_PATH.is_file():
        pytest.skip("Arquivos reais de validação indisponíveis neste computador")
    return EXCEL_PATH.read_bytes(), PDF_PATH.read_bytes()


@pytest.fixture
def sessions(tmp_path: Path):
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'integration.sqlite3'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    yield factory
    engine.dispose()


def _count(session, model) -> int:
    return session.scalar(select(func.count()).select_from(model)) or 0


def _changed_real_workbook(data: bytes) -> tuple[bytes, Decimal]:
    """Change one real service measure in memory to exercise an update."""
    workbook = load_workbook(BytesIO(data))
    try:
        sheet = workbook.active
        assert sheet.cell(1, 1).value == "Serviço"
        assert sheet.cell(1, 10).value == "Cortes"
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row, 1).value) == "20533164":
                updated = Decimal(str(sheet.cell(row, 10).value)) + Decimal("3")
                sheet.cell(row, 10).value = int(updated)
                stream = BytesIO()
                workbook.save(stream)
                return stream.getvalue(), updated
        pytest.fail("Serviço 20533164 ausente da planilha real")
    finally:
        workbook.close()


def test_real_imports_are_linked_idempotent_and_keep_manual_fields(
    real_files: tuple[bytes, bytes], sessions,
) -> None:
    excel_data, pdf_data = real_files

    with sessions.begin() as session:
        excel_result = ImportService(session, None, "ADMIN").import_excel(
            EXCEL_PATH.name, excel_data
        )
        assert excel_result.status == "IMPORTADO"
        assert excel_result.created >= 3

    with sessions.begin() as session:
        pdf_result = ImportService(session, None, "ADMIN").import_pdf(PDF_PATH.name, pdf_data)
        assert pdf_result.status == "VINCULADO"
        assert pdf_result.order_code == "173368"
        assert pdf_result.cart_code == "109633"
        assert pdf_result.items == 9

    with sessions() as session:
        orders = session.scalars(select(Order).where(
            Order.codigo_interno_normalizado == "173368"
        )).all()
        assert len(orders) == 1
        order = orders[0]
        assert order.carrinho == "109633"
        assert order.data_pedido == date(2026, 9, 9)
        assert order.data_carregamento == date(2026, 9, 21)
        assert order.cliente_pdf == "GARBO MOVEIS SOB MEDIDA"
        assert order.valor_produtos == Decimal("3736.31")
        assert order.valor_total == Decimal("3816.21")
        assert order.quantidade_total_pdf == Decimal("681.000")
        assert order.status_dados == "DADOS_COMPLETOS"
        assert order.status_producao == "AGUARDANDO_PROGRAMACAO"
        assert {service.codigo_servico_normalizado for service in order.services} == {"20533164"}
        service = order.services[0]
        assert (service.chapas, service.cortes, service.metros_lineares_corte,
                service.pecas, service.fita_aplicada, service.usinagens) == (
                    Decimal("11"), Decimal("282"), Decimal("227.924"),
                    Decimal("159"), Decimal("120"), Decimal("800"))
        assert len(order.items) == 9
        assert {item.numero_item for item in order.items} == {
            "1", "2", "3", "5", "6", "7", "8", "9", "10"
        }
        two_service_order = session.scalar(select(Order).where(
            Order.codigo_interno_normalizado == "169041"
        ))
        assert two_service_order is not None
        assert {service.codigo_servico_normalizado for service in two_service_order.services} == {
            "23014987", "23041038"
        }
        before = {model: _count(session, model) for model in
                  (Order, Service, OrderItem, Document)}

    with sessions.begin() as session:
        order = session.scalar(select(Order).where(
            Order.codigo_interno_normalizado == "173368"
        ))
        order.prioridade = "URGENTE"
        order.observacao_operacional = "Observação manual preservada"
        order.programado_em = date(2026, 9, 18)

    with sessions.begin() as session:
        importer = ImportService(session, None, "ADMIN")
        assert importer.import_excel(EXCEL_PATH.name, excel_data).status == "DUPLICADO"
        assert importer.import_pdf(PDF_PATH.name, pdf_data).status == "DUPLICADO"

    with sessions.begin() as session:
        importer = ImportService(session, None, "ADMIN")
        importer.import_excel(EXCEL_PATH.name, excel_data, force=True)
        importer.import_pdf(PDF_PATH.name, pdf_data, force=True)

    with sessions() as session:
        assert {model: _count(session, model) for model in before} == before
        order = session.scalar(select(Order).where(
            Order.codigo_interno_normalizado == "173368"
        ))
        assert order.prioridade == "URGENTE"
        assert order.observacao_operacional == "Observação manual preservada"
        assert order.programado_em == date(2026, 9, 18)
        assert len(order.items) == 9
        assert {service.codigo_servico_normalizado for service in order.services} == {"20533164"}


def test_pdf_waits_for_exact_order_and_is_linked_after_excel(
    real_files: tuple[bytes, bytes], sessions,
) -> None:
    excel_data, pdf_data = real_files

    with sessions.begin() as session:
        result = ImportService(session, None, "ADMIN").import_pdf(PDF_PATH.name, pdf_data)
        assert result.status == "PENDENTE"

    with sessions() as session:
        assert _count(session, Order) == 0
        assert _count(session, OrderItem) == 0
        pending = session.scalar(select(PendingLink))
        assert pending is not None
        assert pending.pedido_pdf == "173368"
        assert pending.status == "AGUARDANDO_VINCULACAO"

    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_excel(EXCEL_PATH.name, excel_data)

    with sessions() as session:
        pending = session.scalar(select(PendingLink))
        order = session.scalar(select(Order).where(
            Order.codigo_interno_normalizado == "173368"
        ))
        assert pending.status == "VINCULADO"
        assert order is not None
        assert order.carrinho == "109633"
        assert len(order.items) == 9
        assert _count(session, OrderItem) == 9


def test_changed_service_measure_updates_and_is_audited(
    real_files: tuple[bytes, bytes], sessions,
) -> None:
    excel_data, _ = real_files
    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_excel(EXCEL_PATH.name, excel_data)

    changed_data, new_cuts = _changed_real_workbook(excel_data)
    with sessions.begin() as session:
        result = ImportService(session, None, "ADMIN").import_excel(
            "PLANILHA alterada.xlsx", changed_data
        )
        assert result.updated >= 1

    with sessions() as session:
        order = session.scalar(select(Order).where(
            Order.codigo_interno_normalizado == "173368"
        ))
        service = session.scalar(select(Service).where(
            Service.pedido_id == order.id,
            Service.codigo_servico_normalizado == "20533164"
        ))
        assert service.cortes == new_cuts
        assert len(order.services) == 1
        audits = session.scalars(select(Audit).where(
            Audit.entity == "servicos", Audit.entity_id == str(service.id),
            Audit.field == "cortes", Audit.source == "PLANILHA_SERVICOS"
        )).all()
        assert any(audit.action == "UPDATE" and Decimal(audit.new_value) == new_cuts
                   for audit in audits)


def test_missing_column_preserves_value_but_blank_cell_clears_it(real_files, sessions) -> None:
    excel_data, _ = real_files
    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_excel(EXCEL_PATH.name, excel_data)

    workbook = load_workbook(BytesIO(excel_data))
    try:
        sheet = workbook.active
        sheet.delete_cols(10)
        stream = BytesIO()
        workbook.save(stream)
        partial_data = stream.getvalue()
    finally:
        workbook.close()
    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_excel("sem_cortes.xlsx", partial_data)
    with sessions() as session:
        service = session.scalar(select(Service).where(
            Service.codigo_servico_normalizado == "20533164"))
        assert service.cortes == Decimal("282")

    workbook = load_workbook(BytesIO(excel_data))
    try:
        sheet = workbook.active
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row, 1).value) == "20533164":
                invalid_row = row
                sheet.cell(row, 10).value = "valor inválido"
                break
        stream = BytesIO()
        workbook.save(stream)
        invalid_data = stream.getvalue()
    finally:
        workbook.close()
    with sessions.begin() as session:
        invalid_result = ImportService(session, None, "ADMIN").import_excel(
            "cortes_invalidos.xlsx", invalid_data)
        assert invalid_result.status == "IMPORTADO_COM_ERROS"
    with sessions() as session:
        service = session.scalar(select(Service).where(
            Service.codigo_servico_normalizado == "20533164"))
        assert service.cortes == Decimal("282")
        issue = session.scalar(select(ImportError).where(
            ImportError.documento_id == invalid_result.document_id,
            ImportError.linha == invalid_row, ImportError.severity == "ERRO"))
        assert issue is not None

    workbook = load_workbook(BytesIO(excel_data))
    try:
        sheet = workbook.active
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row, 1).value) == "20533164":
                sheet.cell(row, 10).value = None
                break
        stream = BytesIO()
        workbook.save(stream)
        blank_data = stream.getvalue()
    finally:
        workbook.close()
    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_excel("cortes_vazios.xlsx", blank_data)
    with sessions() as session:
        service = session.scalar(select(Service).where(
            Service.codigo_servico_normalizado == "20533164"))
        assert service.cortes is None
        audit = session.scalar(select(Audit).where(
            Audit.entity_id == str(service.id), Audit.field == "cortes",
            Audit.source == "PLANILHA_SERVICOS", Audit.new_value.is_(None)
        ).order_by(Audit.at.desc()))
        assert audit is not None


def test_revised_pdf_for_same_cart_replaces_current_items(real_files, sessions) -> None:
    excel_data, pdf_data = real_files
    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_excel(EXCEL_PATH.name, excel_data)
    with sessions.begin() as session:
        ImportService(session, None, "ADMIN").import_pdf(PDF_PATH.name, pdf_data)
    with sessions() as session:
        order_count = _count(session, Order)
    pdf = pymupdf.open(stream=pdf_data, filetype="pdf")
    pdf.set_metadata({"title": "nova via do mesmo carrinho"})
    revised = pdf.tobytes()
    pdf.close()
    with sessions.begin() as session:
        result = ImportService(session, None, "ADMIN").import_pdf("segunda_via.pdf", revised)
        assert result.status == "VINCULADO"
    with sessions() as session:
        assert _count(session, Order) == order_count
        assert _count(session, OrderItem) == 9
        pdf_docs = session.scalars(select(Document).where(Document.tipo_documento == "CARRINHO_PDF")).all()
        assert sorted(doc.status for doc in pdf_docs) == ["SUBSTITUIDO", "VINCULADO"]


def test_image_pdf_is_recorded_as_error_without_false_pending(sessions) -> None:
    pdf = pymupdf.open()
    pdf.new_page()
    data = pdf.tobytes()
    pdf.close()
    with sessions.begin() as session:
        result = ImportService(session, None, "ADMIN").import_pdf("scan.pdf", data)
        assert result.status == "ERRO_LEITURA"
        assert "sem texto selecionável" in result.warnings[0]
    with sessions() as session:
        assert _count(session, Document) == 1
        assert session.scalar(select(Document)).status == "ERRO_IMPORTACAO"
        assert _count(session, ImportError) == 1
        assert _count(session, PendingLink) == 0
    with sessions.begin() as session:
        duplicate = ImportService(session, None, "ADMIN").import_pdf("scan.pdf", data)
        assert duplicate.status == "DUPLICADO"
        retry = ImportService(session, None, "ADMIN").import_pdf("scan.pdf", data, force=True)
        assert retry.status == "ERRO_LEITURA"
    with sessions() as session:
        assert _count(session, Document) == 1
        assert _count(session, ImportError) == 1
        assert _count(session, PendingLink) == 0
