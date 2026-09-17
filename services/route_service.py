"""Route spreadsheet lookup for delivery-date validation."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook


ROUTE_FILE = Path(__file__).resolve().parents[1] / "ROTAS_2026.xlsx"
WEEKDAYS = ("SEGUNDA", "TERCA", "QUARTA", "QUINTA", "SEXTA")
STATE_UFS = {
    "AC", "AL", "AM", "AP", "BA", "CE", "DF", "ES", "GO", "MA", "MG", "MS", "MT",
    "PA", "PB", "PE", "PI", "PR", "RJ", "RN", "RO", "RR", "RS", "SC", "SE", "SP",
    "TO",
}


@dataclass(frozen=True)
class RouteEntry:
    weekday: int
    city: str
    route: str
    source: str


@dataclass(frozen=True)
class RouteCheck:
    ok: bool
    city: str
    delivery_date: date
    weekday: str
    routes: tuple[str, ...]
    message: str


def normalize_city(value: object) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"\([^)]*\)", " ", text)
    text = re.sub(r"\bR\s*\.?\s*\d+\b", " ", text, flags=re.I)
    text = text.replace("/", " ")
    text = text.replace("-", " ")
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).strip().upper()
    text = re.sub(r"\bCONDICAO\b", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    parts = text.split()
    if len(parts) > 1 and parts[-1] in STATE_UFS:
        text = " ".join(parts[:-1])
    return text


def _city_keys(value: object) -> set[str]:
    base = normalize_city(value)
    if not base:
        return set()
    keys = {base}
    for marker in (" CONDICAO", " REGIAO"):
        if marker in base:
            keys.add(base.split(marker, 1)[0].strip())
    if " " in base:
        keys.add(base.replace(" ", ""))
    return {key for key in keys if key}


def _weekday_from_header(value: object, fallback: int) -> int:
    if isinstance(value, datetime):
        return value.date().weekday()
    if isinstance(value, date):
        return value.weekday()
    text = normalize_city(value)
    for index, label in enumerate(WEEKDAYS):
        if label in text:
            return index
    return fallback


def _looks_like_route(value: str) -> bool:
    return bool(re.search(r"\(\s*R\s*\.?\s*\d+", value, re.I))


@lru_cache(maxsize=4)
def route_entries(path: str | None = None) -> tuple[RouteEntry, ...]:
    source_path = Path(path) if path else ROUTE_FILE
    if not source_path.exists():
        return tuple()
    workbook = load_workbook(source_path, read_only=True, data_only=True)
    try:
        sheet = workbook["CIDADES X ROTAS ATUALIZADAS"]
        entries: list[RouteEntry] = []
        for column in range(1, min(sheet.max_column, 5) + 1):
            weekday = _weekday_from_header(sheet.cell(1, column).value, column - 1)
            current_route = ""
            for row in range(2, sheet.max_row + 1):
                value = sheet.cell(row, column).value
                if not isinstance(value, str) or not value.strip():
                    continue
                label = value.strip()
                if _looks_like_route(label):
                    current_route = label
                for key in _city_keys(label):
                    entries.append(RouteEntry(weekday, key, current_route or label, str(source_path)))
        return tuple(entries)
    finally:
        workbook.close()


def valid_routes_for(city: object, delivery_date: date) -> tuple[str, ...]:
    keys = _city_keys(city)
    if not keys or delivery_date.weekday() > 4:
        return tuple()
    routes = {
        entry.route
        for entry in route_entries()
        if entry.weekday == delivery_date.weekday() and entry.city in keys
    }
    return tuple(sorted(routes))


def allowed_weekdays_for(city: object) -> tuple[str, ...]:
    keys = _city_keys(city)
    indexes = sorted({entry.weekday for entry in route_entries() if entry.city in keys})
    return tuple(WEEKDAYS[index] for index in indexes if 0 <= index < len(WEEKDAYS))


def check_delivery_route(city: object, delivery_date: date) -> RouteCheck:
    city_text = "" if city is None else str(city).strip()
    weekday = WEEKDAYS[delivery_date.weekday()] if delivery_date.weekday() < 5 else "FIM DE SEMANA"
    if not city_text:
        return RouteCheck(False, city_text, delivery_date, weekday, tuple(), "Cidade do pedido ainda nao foi identificada.")
    routes = valid_routes_for(city_text, delivery_date)
    if routes:
        return RouteCheck(True, city_text, delivery_date, weekday, routes, f"{city_text} atende rota de {weekday}.")
    allowed = allowed_weekdays_for(city_text)
    if allowed:
        return RouteCheck(False, city_text, delivery_date, weekday, tuple(),
                          f"{city_text} nao consta na rota de {weekday}. Dias encontrados: {', '.join(allowed)}.")
    return RouteCheck(False, city_text, delivery_date, weekday, tuple(),
                      f"{city_text} nao foi encontrada na planilha de rotas.")
